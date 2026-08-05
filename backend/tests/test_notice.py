"""The hearing notice — quoted verbatim, as the settled ¶4 does.

    4. That despite the earlier inclusion ... she received a hearing notice
       bearing Notice No.: EFS... on 30.12.2025, wherein, inter alia, it was
       stated that:

            "after scrutiny it has been observed that: ..."

       Vide the said Notice, [the Appellant] was asked to appear before
       the AERO, Ward 42 ...
"""
from __future__ import annotations

from docx.shared import Inches

from app.generator import build_petition
from app.intake import PetitionInput
from tests.helpers import all_text

NOTICE_TEXT = (
    "after scrutiny it has been observed that: You have not filled in the "
    "details related to you or your relative that could establish you or your "
    "relative as a registered elector in the electoral rolls prepared during "
    "the previous SIR."
)


def minimal_intake(**overrides) -> PetitionInput:
    data = dict(appellant_name="Rehana Bibi Mondal", gender="female")
    data.update(overrides)
    return PetitionInput(**data)


def noticed(**overrides) -> PetitionInput:
    data = dict(
        notice_received=True,
        notice_no="EFS0000000000000000000001",
        notice_date="30.12.2025",
    )
    data.update(overrides)
    return minimal_intake(**data)


def test_notice_reason_is_quoted_in_an_indented_block():
    doc = build_petition(noticed(notice_reason=NOTICE_TEXT))
    quote = next(p for p in doc.paragraphs if "after scrutiny" in p.text)
    assert quote.text.strip().startswith('"')
    assert quote.text.strip().endswith('"')
    assert quote.paragraph_format.left_indent == Inches(0.75)
    assert "wherein, inter alia, it was stated that:" in all_text(doc)


def test_notice_without_a_quoted_reason_carries_no_block():
    doc = build_petition(noticed())
    assert not any(p.text.strip().startswith('"') for p in doc.paragraphs)
    assert "inter alia" not in all_text(doc)


def test_vide_line_names_the_officer_to_appear_before():
    doc = build_petition(
        noticed(hearing_officer="the AERO, Ward 42")
    )
    vide = next(p for p in doc.paragraphs if "Vide the said Notice" in p.text)
    assert "the AERO, Ward 42" in vide.text
    # Third person, though the settled draft slips into "I was asked".
    assert "the Appellant was asked to appear" in vide.text
    assert vide.paragraph_format.left_indent == Inches(0.5)


def test_vide_line_absent_when_no_officer_is_named():
    assert "Vide the said Notice" not in all_text(build_petition(noticed()))
